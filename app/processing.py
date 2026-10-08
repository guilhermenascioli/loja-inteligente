"""Regras de validação e cálculos; sem rede, persistência ou IA."""
import csv
import io
import re
import unicodedata
import zipfile
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

SCHEMAS = {
    'vendas': ['data', 'produto', 'categoria', 'quantidade', 'preco_unitario'],
    'despesas': ['data', 'descricao', 'categoria', 'valor'],
}
MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 20000


class FileError(ValueError):
    pass


def normalize(value):
    text = unicodedata.normalize('NFKD', str(value or '').strip().lower())
    text = ''.join(c for c in text if not unicodedata.combining(c))
    return re.sub(r'\s+', '_', text)


def number(value, money=False):
    if isinstance(value, bool) or value is None:
        raise ValueError('Informe um número válido.')
    text = str(value).strip()
    if text.startswith('='):
        raise ValueError('Substitua a fórmula pelo valor calculado.')
    text = text.replace('R$', '').replace(' ', '')
    if ',' in text:
        if not re.fullmatch(r'[+-]?(?:\d+|\d{1,3}(?:\.\d{3})+),\d+', text):
            raise ValueError('Número inválido. Use 1234,56 ou 1234.56.')
        text = text.replace('.', '').replace(',', '.')
    elif not re.fullmatch(r'[+-]?\d+(?:\.\d+)?', text):
        raise ValueError('Número inválido. Use 1234,56 ou 1234.56.')
    try:
        result = Decimal(text)
    except InvalidOperation:
        raise ValueError('Informe um número válido.') from None
    if not result.is_finite() or abs(result) > Decimal('1000000000'):
        raise ValueError('Número fora do limite de 1 bilhão.')
    if money and result != result.quantize(Decimal('.01')):
        raise ValueError('Valores em dinheiro devem ter no máximo duas casas decimais.')
    return result


def parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for fmt in ('%d/%m/%Y', '%Y-%m-%d'):
        try:
            return datetime.strptime(str(value).strip(), fmt).date()
        except ValueError:
            pass
    raise ValueError('Data inválida. Use DD/MM/AAAA ou AAAA-MM-DD.')


def read_rows(content, filename, sheet_name=None):
    if not content:
        raise FileError('O arquivo está vazio.')
    if len(content) > MAX_BYTES:
        raise FileError('O limite é 5 MB por arquivo.')
    suffix = filename.rsplit('.', 1)[-1].lower()
    notes = []
    try:
        if suffix == 'xlsx':
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                    raise FileError('A planilha descompactada excede o limite de 50 MB.')
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
            try:
                if sheet_name and sheet_name not in book.sheetnames:
                    raise FileError('A aba selecionada não existe.')
                sheet = book[sheet_name] if sheet_name else book.worksheets[0]
                if len(book.worksheets) > 1:
                    notes.append('Aba analisada: ' + sheet.title + '. As demais abas não entram nos indicadores.')
                if sheet.max_column and sheet.max_column > 100:
                    raise FileError('A planilha deve ter no máximo 100 colunas.')
                rows = []
                for row in sheet.iter_rows(values_only=True):
                    rows.append(list(row))
                    if len(rows) > MAX_ROWS + 1:
                        raise FileError('O limite é 20.000 registros por arquivo.')
            finally:
                book.close()
        elif suffix == 'csv':
            try:
                text = content.decode('utf-8-sig')
            except UnicodeDecodeError:
                text = content.decode('cp1252')
                notes.append('Arquivo lido com codificação Windows (CP1252).')
            sample = text.splitlines()[0] if text.splitlines() else ''
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=',;\t')
            except csv.Error:
                dialect = csv.excel
            rows = []
            for row in csv.reader(io.StringIO(text), dialect, strict=True):
                if len(row) > 100:
                    raise FileError('O arquivo deve ter no máximo 100 colunas.')
                rows.append(row)
                if len(rows) > MAX_ROWS + 1:
                    raise FileError('O limite é 20.000 registros por arquivo.')
        else:
            raise FileError('Escolha um arquivo XLSX ou CSV.')
    except FileError:
        raise
    except Exception:
        raise FileError('Não foi possível ler o arquivo. Verifique se é um XLSX ou CSV válido, sem senha.') from None
    if len(rows) < 2:
        raise FileError('Inclua os cabeçalhos na primeira linha e pelo menos um registro abaixo.')
    return rows, notes


def cash(value):
    return float(value.quantize(Decimal('.01'), rounding=ROUND_HALF_UP))


def analyze(content, filename, kind, mapping=None, sheet_name=None):
    if kind not in SCHEMAS:
        raise FileError('Escolha vendas ou despesas.')
    rows, notes = read_rows(content, filename, sheet_name)
    headers = [normalize(c) for c in rows[0]]
    if len(set(headers)) != len(headers) or '' in headers:
        raise FileError('Há cabeçalhos vazios ou repetidos. Corrija a primeira linha.')
    if mapping is not None:
        if not isinstance(mapping, dict) or set(mapping) != set(SCHEMAS[kind]):
            raise FileError('Associe todas as colunas necessárias.')
        selected = list(mapping.values())
        if any(not isinstance(v, str) or v not in headers for v in selected) or len(set(selected)) != len(selected):
            raise FileError('Selecione uma coluna diferente e existente para cada campo.')
    else:
        mapping = {field: field for field in SCHEMAS[kind]}
    missing = [c for c in SCHEMAS[kind] if mapping[c] not in headers]
    if missing:
        raise FileError('Colunas necessárias ausentes: ' + ', '.join(missing) + '.')
    extras = [c for c in headers if c not in mapping.values()]
    if extras:
        notes.append('Colunas adicionais preservadas na aba Originais e não utilizadas nos cálculos: ' + ', '.join(extras) + '.')
    records, issues, originals, seen = [], [], [], {}
    for line, values in enumerate(rows[1:], 2):
        raw = {h: str(values[i]) if i < len(values) and values[i] is not None else '' for i, h in enumerate(headers)}
        originals.append({'linha': line, **raw})
        row = {h: values[i] if i < len(values) else None for i, h in enumerate(headers)}
        errors = []
        clean = {'linha': line}
        if len(values) > len(headers) and any(v not in (None, '') for v in values[len(headers):]):
            errors.append(('estrutura', 'Existem valores sem cabeçalho nesta linha.'))
            originals[-1]['valores_sem_cabecalho'] = str(values[len(headers):])
        for field in SCHEMAS[kind]:
            value = row.get(mapping[field])
            try:
                if value is None or str(value).strip() == '':
                    raise ValueError('Campo obrigatório ausente.')
                if field == 'data':
                    clean[field] = parse_date(value).isoformat()
                elif field in ('quantidade', 'preco_unitario', 'valor'):
                    n = number(value, money=field != 'quantidade')
                    if n < 0 or (field == 'quantidade' and (n == 0 or n != n.to_integral_value())):
                        raise ValueError('Use quantidade inteira positiva.' if field == 'quantidade' else 'Use valor maior ou igual a zero. Estornos não são suportados nesta versão.')
                    clean[field] = n
                else:
                    if str(value).strip().startswith('='):
                        raise ValueError('Substitua a fórmula pelo valor calculado.')
                    clean[field] = str(value).strip()
            except ValueError as e:
                errors.append((field, str(e)))
        if errors:
            for field, message in errors:
                issues.append({'linha': line, 'campo': field, 'nivel': 'erro', 'mensagem': message})
            originals[-1]['situacao'] = 'Fora dos indicadores'
            continue
        key = tuple(clean[c] for c in SCHEMAS[kind])
        duplicate = key in seen
        if duplicate:
            issues.append({'linha': line, 'campo': 'registro', 'nivel': 'aviso', 'mensagem': f'Registro igual à linha {seen[key]}. Mantido nos cálculos; confirme se é uma operação distinta.'})
        else:
            seen[key] = line
        total = clean['quantidade'] * clean['preco_unitario'] if kind == 'vendas' else clean['valor']
        clean['total'] = total.quantize(Decimal('.01'), rounding=ROUND_HALF_UP)
        clean['duplicado'] = duplicate
        originals[-1]['situacao'] = 'Incluído; possível duplicado' if duplicate else 'Incluído'
        records.append(clean)
    monthly, groups = defaultdict(Decimal), defaultdict(Decimal)
    for r in records:
        monthly[r['data'][:7]] += r['total']
        groups[r['produto'] if kind == 'vendas' else r['categoria']] += r['total']
    if records:
        start, end = min(monthly), max(monthly)
        year, month = map(int, start.split('-'))
        while f'{year:04}-{month:02}' <= end:
            monthly[f'{year:04}-{month:02}'] += Decimal(0)
            year, month = (year + 1, 1) if month == 12 else (year, month + 1)
    total = sum((r['total'] for r in records), Decimal(0))
    ranking = [{'nome': name, 'total': cash(amount), 'participacao': round(float(amount / total * 100), 2) if total else 0} for name, amount in sorted(groups.items(), key=lambda p: (-p[1], p[0]))]
    result = {
        'tipo': kind, 'total': cash(total),
        'quantidade': int(sum(r['quantidade'] for r in records)) if kind == 'vendas' else None,
        'recebidos': len(rows) - 1, 'validos': len(records), 'invalidos': len(rows) - 1 - len(records),
        'duplicados': sum(r['duplicado'] for r in records),
        'inicio': min((r['data'] for r in records), default=None), 'fim': max((r['data'] for r in records), default=None),
        'mensal': [{'mes': m, 'total': cash(v)} for m, v in sorted(monthly.items())],
        'ranking': ranking, 'maiores': [{'data': r['data'], 'descricao': r.get('descricao', r.get('produto')), 'total': cash(r['total'])} for r in sorted(records, key=lambda r: -r['total'])[:10]],
        'problemas': issues, 'avisos': notes,
    }
    # JSON/export never receive Decimal objects.
    result['tratados'] = [{k: cash(v) if isinstance(v, Decimal) else v for k, v in r.items()} for r in records]
    result['originais'] = originals
    return result


def safe_cell(value):
    # Preserve text while preventing Excel formula injection, including raw uploads.
    if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')):
        return "'" + value
    return value


def workbook_bytes(tables):
    book = Workbook()
    book.remove(book.active)
    for name, rows in tables.items():
        sheet = book.create_sheet(name)
        headers = list(dict.fromkeys(key for row in rows for key in row)) if rows else ['informacao']
        sheet.append(headers)
        for row in rows:
            sheet.append([safe_cell(row.get(h, '')) for h in headers])
        if not rows:
            sheet.append(['Nenhum registro nesta seção.'])
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for c in sheet[1]:
            c.fill = PatternFill('solid', fgColor='145C41')
            c.font = Font(color='FFFFFF', bold=True)
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(65, max(16, max(len(str(c.value or '')) for c in column) + 2))
            for cell in column[1:]:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                if isinstance(cell.value, float):
                    cell.number_format = '#,##0.00'
        sheet.sheet_view.showGridLines = False
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def export(result):
    return workbook_bytes({
        'Indicadores': [{'indicador': k, 'valor': result[k]} for k in ('tipo', 'total', 'quantidade', 'recebidos', 'validos', 'invalidos', 'duplicados', 'inicio', 'fim')] + [{'indicador': 'Regra', 'valor': 'Duplicados mantidos. Linhas inválidas fora dos indicadores. Dinheiro arredondado por operação; quantidade inteira positiva.'}] + [{'indicador': 'Aviso', 'valor': n} for n in result['avisos']],
        'Mensal': result['mensal'], 'Participacao': result['ranking'], 'Maiores operacoes': result['maiores'],
        'Tratados': result['tratados'], 'Problemas': result['problemas'], 'Originais': result['originais'],
    })


def example(kind):
    if kind not in SCHEMAS:
        raise FileError('Escolha vendas ou despesas.')
    rows = []
    for month in range(1, 7):
        if kind == 'vendas':
            for i, (product, category, price) in enumerate([('Café 500g', 'Mercearia', 18.90), ('Sabonete', 'Higiene', 4.50), ('Arroz 5kg', 'Mercearia', 29.90), ('Suco 1L', 'Bebidas', 8.90)]):
                rows.append(dict(zip(SCHEMAS[kind], [f'2026-{month:02}-{(i+1)*4:02}', product, category, 12 + month * 3 + i * 2, price])))
        else:
            for i, (desc, category, value) in enumerate([('Aluguel da loja', 'Estrutura', 1100), ('Conta de energia', 'Serviços', 190 + month * 7), ('Entrega de mercadorias', 'Transporte', 80 + month * 3), ('Anúncios locais', 'Marketing', 120 + month * 5)]):
                rows.append(dict(zip(SCHEMAS[kind], [f'2026-{month:02}-{(i+1)*5:02}', desc, category, value])))
    return workbook_bytes({'Vendas' if kind == 'vendas' else 'Despesas': rows})


def inspect_file(content, filename, kind, sheet_name=None):
    if kind not in SCHEMAS:
        raise FileError('Escolha vendas ou despesas.')
    rows, notes = read_rows(content, filename, sheet_name)
    headers = [normalize(c) for c in rows[0]]
    if len(set(headers)) != len(headers) or '' in headers:
        raise FileError('Há cabeçalhos vazios ou repetidos. Corrija a primeira linha.')
    sheets = []
    if filename.lower().endswith('.xlsx'):
        book = load_workbook(io.BytesIO(content), read_only=True, keep_links=False)
        try:
            sheets = book.sheetnames
        finally:
            book.close()
    aliases = {'data': ['dia', 'dia_da_venda', 'data_venda', 'data_da_venda'],
               'produto': ['item', 'nome_do_item', 'nome_do_produto'],
               'categoria': ['grupo', 'tipo', 'categoria_do_produto'],
               'quantidade': ['qtd', 'qtde', 'unidades'],
               'preco_unitario': ['preco', 'valor_unitario'],
               'descricao': ['historico', 'despesa', 'descrição_da_despesa'],
               'valor': ['valor_pago', 'valor_da_despesa', 'total']}
    suggestions = {}
    for field in SCHEMAS[kind]:
        matches = [h for h in headers if h == field or h in aliases.get(field, [])]
        suggestions[field] = matches[0] if len(matches) == 1 else ''
    return {'colunas': [{'nome': str(raw), 'chave': h} for raw, h in zip(rows[0], headers)],
            'abas': sheets, 'aba': sheet_name or (sheets[0] if sheets else ''),
            'campos': SCHEMAS[kind], 'sugestoes': suggestions, 'registros': len(rows)-1}
