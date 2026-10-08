import io
import json
import threading
import unittest
import urllib.error
import urllib.request
from http.server import ThreadingHTTPServer
from unittest.mock import patch
from openpyxl import load_workbook
from app.server import Handler
from app.processing import example


class ServerTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.server = ThreadingHTTPServer(('127.0.0.1',0), Handler)
        cls.base = f'http://127.0.0.1:{cls.server.server_port}'
        cls.thread = threading.Thread(target=cls.server.serve_forever,daemon=True)
        cls.thread.start()
        cls.opener = urllib.request.OpenerDirector()
        for handler in (urllib.request.HTTPHandler(), urllib.request.HTTPDefaultErrorHandler(), urllib.request.HTTPErrorProcessor()):
            cls.opener.add_handler(handler)

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown()
        cls.server.server_close()
        cls.thread.join(timeout=3)

    def request(self,path,body=None,headers=None):
        return self.opener.open(urllib.request.Request(self.base+path,data=body,headers=headers or {}),timeout=5)

    def test_index_assets_status_and_hidden_files(self):
        for path in ('/','/app.js','/style.css','/api/status'):
            with self.request(path) as response:
                self.assertEqual(response.status,200)
                self.assertEqual(response.headers['Cache-Control'],'no-store')
        for path in ('/.env','/app/ai.py','/../README.md'):
            with self.assertRaises(urllib.error.HTTPError) as caught:self.request(path)
            self.assertEqual(caught.exception.code,404)

    def test_example_analyze_export(self):
        with self.request('/api/example?tipo=vendas') as response:content=response.read()
        route='?tipo=vendas&nome=exemplo.xlsx'
        with self.request('/api/analyze'+route,content) as response:r=json.load(response)
        self.assertEqual(r['total'],9489)
        self.assertNotIn('originais',r)
        self.assertNotIn('tratados',r)
        with self.request('/api/export'+route,content) as response:
            book=load_workbook(io.BytesIO(response.read()))
            self.assertIn('Problemas',book.sheetnames)

    def test_origin_and_host_rejected(self):
        for headers in ({'Origin':'https://example.com'},{'Host':'example.com'}):
            with self.assertRaises(urllib.error.HTTPError) as caught:self.request('/api/status',headers=headers)
            self.assertEqual(caught.exception.code,403)

    def test_invalid_file_and_consent_before_ai(self):
        with self.assertRaises(urllib.error.HTTPError) as caught:self.request('/api/analyze?tipo=vendas&nome=x.csv',b'bad')
        self.assertEqual(caught.exception.code,400)
        with patch('app.ai.explain') as call:
            with self.assertRaises(urllib.error.HTTPError) as caught:self.request('/api/explain?tipo=vendas&nome=x.xlsx',example('vendas'))
            self.assertEqual(caught.exception.code,400)
            call.assert_not_called()

    def test_explain_uses_server_results(self):
        with patch('app.ai.explain',return_value='Resumo simulado para o teste.') as call:
            with self.request('/api/explain?tipo=despesas&nome=x.xlsx',example('despesas'),{'X-Loja-Consent':'sim'}) as response:
                self.assertEqual(json.load(response)['texto'],'Resumo simulado para o teste.')
            self.assertEqual(call.call_args.args[0]['total'],9255)


if __name__=='__main__':
    unittest.main()
