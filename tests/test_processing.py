import io
import json
import os
import unittest
from unittest.mock import patch
from openpyxl import Workbook, load_workbook
from app.processing import FileError, analyze, example, export, number
from app import ai


def csv(text):
    return text.encode('utf-8')


class ProcessingTests(unittest.TestCase):
    def sales(self, rows):
        return analyze(csv('data;produto;categoria;quantidade;preço unitário\n'+rows),'vendas.csv','vendas')

    def test_sales_decimal_months_and_participation(self):
        r=self.sales('01/01/2026;Café;Alimentos;3;0,10\n2026-03-01;Arroz;Alimentos;2;10,00')
        self.assertEqual(r['total'],20.30)
        self.assertEqual(r['quantidade'],5)
        self.assertEqual(r['mensal'],[{'mes':'2026-01','total':0.3},{'mes':'2026-02','total':0.0},{'mes':'2026-03','total':20.0}])
        self.assertEqual(r['ranking'][0]['nome'],'Arroz')
        self.assertAlmostEqual(sum(x['participacao'] for x in r['ranking']),100,places=2)

    def test_invalid_dates_amounts_missing_and_duplicates_visible(self):
        r=self.sales('01/01/2026;Café;Alimentos;2;10\n01/01/2026;Café;Alimentos;2;10\n31/02/2026;X;Y;2;5\n01/01/2026;;Y;2;5\n01/01/2026;X;Y;1,5;5\n01/01/2026;X;Y;2;-1\n;;;;')
        self.assertEqual((r['recebidos'],r['validos'],r['invalidos'],r['duplicados']),(7,2,5,1))
        self.assertEqual(r['total'],40)
        self.assertEqual(len(r['originais']),7)
        self.assertTrue(any(x['campo']=='data' for x in r['problemas']))

    def test_locale_numbers_and_nonfinite(self):
        self.assertEqual(str(number('R$ 1.234,56',True)),'1234.56')
        for value in ('NaN','inf','1,2,3','1.234,5.6','1e9',True,'=2+2'):
            with self.subTest(value=value),self.assertRaises(ValueError):number(value)
        with self.assertRaises(ValueError):number('1.234',True)

    def test_expenses_categories_and_largest(self):
        r=analyze(csv('data,descricao,categoria,valor\n2026-01-01,Aluguel,Estrutura,1000\n2026-01-02,Luz,Servicos,150\n2026-02-01,Luz,Servicos,200'),'gastos.csv','despesas')
        self.assertEqual(r['total'],1350)
        self.assertEqual(r['ranking'][0]['total'],1000)
        self.assertEqual(r['maiores'][0]['descricao'],'Aluguel')

    def test_missing_empty_corrupt_and_wrong_type(self):
        for content,name in [(b'','x.csv'),(b'a,b\n1,2','x.csv'),(b'bad','x.xlsx'),(b'abc','x.txt')]:
            with self.subTest(name=name),self.assertRaises(FileError):analyze(content,name,'vendas')

    def test_duplicate_headers_and_extra_row_fields(self):
        with self.assertRaises(FileError):self.sales('')
        with self.assertRaises(FileError):analyze(b'data,data\n1,2','x.csv','vendas')
        r=self.sales('01/01/2026;X;Y;1;2;extra')
        self.assertEqual(r['invalidos'],1)
        self.assertIn('extra',r['originais'][0]['valores_sem_cabecalho'])

    def test_zero_valid_and_zero_value(self):
        r=self.sales('bad;X;Y;1;2')
        self.assertEqual(r['validos'],0)
        self.assertEqual(r['ranking'],[])
        r=self.sales('2026-01-01;X;Y;1;0')
        self.assertEqual(r['validos'],1)
        self.assertEqual(r['ranking'][0]['participacao'],0)

    def test_xlsx_dates_formulas_and_multiple_sheets(self):
        from datetime import datetime
        book=Workbook();sheet=book.active
        sheet.append(['data','produto','categoria','quantidade','preco_unitario'])
        sheet.append([datetime(2026,1,1),'X','Y',2,10]);sheet.append([datetime(2026,1,1),'X','Y',2,'=10+5'])
        book.create_sheet('Ignorada');buffer=io.BytesIO();book.save(buffer)
        r=analyze(buffer.getvalue(),'a.xlsx','vendas')
        self.assertEqual((r['validos'],r['invalidos']),(1,1))
        self.assertTrue(r['avisos'])

    def test_examples_export_and_formula_injection(self):
        for kind,total in [('vendas',9489.0),('despesas',9255.0)]:
            r=analyze(example(kind),'x.xlsx',kind)
            self.assertEqual(r['validos'],24)
            self.assertEqual(r['total'],total)
            if kind=='vendas':self.assertEqual(r['quantidade'],612)
            book=load_workbook(io.BytesIO(export(r)))
            self.assertIn('Originais',book.sheetnames)
            self.assertEqual(book['Tratados'].max_row,25)
            self.assertEqual(book['Indicadores']['B3'].value,r['total'])
        r=self.sales('2026-01-01;@texto;Y;1;2\n2026-01-01;X;Y;1;=2+2')
        book=load_workbook(io.BytesIO(export(r)))
        for sheet in book:
            self.assertFalse(any(c.data_type=='f' for row in sheet for c in row))

    def test_limits(self):
        with self.assertRaises(FileError):analyze(b'x'*(5*1024*1024+1),'x.csv','vendas')
        with self.assertRaises(FileError):self.sales('2026-01-01;X;Y;1;1\n'*20001)


class AITests(unittest.TestCase):
    def result(self):
        return analyze(example('vendas'),'a.xlsx','vendas')

    def test_no_configuration(self):
        with patch.dict(os.environ,{},clear=True),self.assertRaises(ai.AIError):ai.explain(self.result())

    def test_payload_no_names_or_raw_records(self):
        content=json.dumps(ai.payload(self.result()))
        for private in ('Café','Sabonete','Mercearia','originais','tratados','descricao'):
            self.assertNotIn(private,content)

    def test_api_failure_sanitized(self):
        import urllib.error
        with patch.dict(os.environ,{'LOJA_OLLAMA_MODEL':'test-model'},clear=True),patch('app.ai.api_request',side_effect=urllib.error.URLError('test-secret')):
            with self.assertRaises(ai.AIError) as caught:ai.explain(self.result())
            self.assertNotIn('test-secret',str(caught.exception))

    def test_success_and_numeric_response_rejected(self):
        for text,valid in [('As vendas aumentaram ao longo do período.',True),('O faturamento foi 999.',False)]:
            response={'done':True,'response':text}
            with patch.dict(os.environ,{'LOJA_OLLAMA_MODEL':'test'},clear=True),patch('app.ai.api_request',side_effect=[{},response]):
                if valid:self.assertEqual(ai.explain(self.result()),text)
                else:
                    with self.assertRaises(ai.AIError):ai.explain(self.result())

    def test_gemini_requires_explicit_free_confirmation(self):
        with patch.dict(os.environ,{'LOJA_GEMINI_API_KEY':'test','LOJA_GEMINI_MODEL':'gemini-test'},clear=True):
            self.assertFalse(ai.configured())

    def test_gemini_body_and_response(self):
        response={'candidates':[{'finishReason':'STOP','content':{'parts':[{'text':'As vendas aumentaram no período.'}]}}]}
        with patch.dict(os.environ,{'LOJA_GEMINI_API_KEY':'test','LOJA_GEMINI_MODEL':'gemini-test','LOJA_GEMINI_FREE_TIER_CONFIRMED':'sim'},clear=True),patch('app.ai.api_request',return_value=response) as request:
            self.assertEqual(ai.explain(self.result()),'As vendas aumentaram no período.')
            self.assertNotIn('Café',json.dumps(request.call_args.args[1]))
            self.assertEqual(request.call_args.args[2]['x-goog-api-key'],'test')

    def test_ollama_cloud_rejected(self):
        with patch.dict(os.environ,{'LOJA_OLLAMA_MODEL':'model-cloud'},clear=True):
            self.assertFalse(ai.configured())
        with patch.dict(os.environ,{'LOJA_OLLAMA_MODEL':'model'},clear=True),patch('app.ai.api_request',return_value={'remote_host':'external'}),self.assertRaises(ai.AIError):
            ai.explain(self.result())


if __name__=='__main__':
    unittest.main()
